import asyncio
import io
import math
import re
from collections import defaultdict
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy.ext.asyncio import AsyncSession

from myapp.schemas.product_passports import (
    ProductPassportNodeResponse,
    ProductPassportResponse,
)
from myapp.services.product_passport_service import ProductPassportService


class ProductPassportExportService:
    """Выгрузка сохранённого паспорта"""

    HEADERS = (
        "Состав паспорта",
        "Заводской номер",
        "Обозначение",
        "Артикул",
        "Полное наименование",
        "Изготовитель",
        "Поставщик",
        "Дата изготовления",
        "Дата установки",
        "Ввод в эксплуатацию",
    )
    WIDTHS = (82, 20, 32, 19, 50, 35, 35, 20, 20, 23)

    @staticmethod
    def _section_order(node: ProductPassportNodeResponse) -> int:
        for name in (node.tree_name, node.full_name):
            normalized = (name or "").strip().upper()
            if "БУСТЕР" in normalized:
                return 2
            section = re.match(
                r"""^СЕКЦИЯ\s+[«"']?([АAБB])(?=$|[\s»"'.,:№-])""", normalized
            )
            if section:
                return 0 if section[1] in ("А", "A") else 1
        return 3

    @classmethod
    def _ordered_nodes(cls, passport: ProductPassportResponse):
        children = defaultdict(list)
        for node in passport.nodes:
            children[node.parent_id].append(node)
        roots = children[None]
        if len(roots) != 1:
            raise ValueError("У паспорта должен быть один корневой узел")
        children[roots[0].id].sort(key=cls._section_order)
        stack = [(roots[0], 0, "", True)]
        visited = set()
        ordered = []
        while stack:
            node, depth, prefix, last = stack.pop()
            if node.id in visited:
                raise ValueError("Повторный узел в дереве паспорта")
            visited.add(node.id)
            branch = children[node.id]
            label = (
                prefix + ("└─ " if last else "├─ ") + node.tree_name
                if depth
                else node.tree_name
            )
            ordered.append((node, depth, label, bool(branch)))
            next_prefix = prefix + ("   " if last else "│  ") if depth else ""
            for index in range(len(branch) - 1, -1, -1):
                stack.append(
                    (branch[index], depth + 1, next_prefix, index == len(branch) - 1)
                )
        if len(visited) != len(passport.nodes):
            raise ValueError("Есть узлы вне дерева паспорта")
        return ordered

    @staticmethod
    def _text(cell, value):
        cell.value = value
        # Названия и номера из источника всегда являются текстом, включая строки с '='.
        if isinstance(value, str):
            cell.data_type = "s"

    @classmethod
    def generate_excel(cls, passport: ProductPassportResponse) -> io.BytesIO:
        ordered = cls._ordered_nodes(passport)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Паспорт"
        sheet.sheet_view.showGridLines = False
        sheet.sheet_view.zoomScale = 85
        sheet.freeze_panes = "B6"
        sheet.sheet_properties.outlinePr.summaryBelow = False
        sheet.sheet_properties.outlinePr.showOutlineSymbols = True
        sheet.sheet_format.outlineLevelRow = min(
            max(depth for _, depth, _, _ in ordered), 7
        )
        for column, width in enumerate(cls.WIDTHS, 1):
            sheet.column_dimensions[sheet.cell(5, column).column_letter].width = width
        cls._text(sheet["A2"], passport.omega_name)
        sheet["A2"].font = Font(name="Arial", size=16, bold=True, color="172B4D")
        sheet.row_dimensions[2].height = 32
        cls._text(
            sheet["A3"],
            f"Источник: сохранённый паспорт Omega. Выгрузка {datetime.now():%d.%m.%Y}.",
        )
        sheet["A3"].font = Font(name="Arial", size=10, color="64748B")
        cls._text(sheet["B3"], "Ветки: + / − слева")
        sheet["B3"].font = Font(name="Arial", size=10, color="64748B")
        sheet.row_dimensions[3].height = 26
        sheet.row_dimensions[5].height = 36
        for column, title in enumerate(cls.HEADERS, 1):
            cell = sheet.cell(5, column, title)
            cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="334867")
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            cell.border = Border(right=Side(style="thin", color="FFFFFF"))

        capacities = (76, 18, 29, 17, 46, 32, 32, 18, 18, 21)
        for index, (node, depth, label, has_children) in enumerate(ordered):
            row = index + 6
            values = (
                label,
                node.serial_number,
                node.designation,
                node.article,
                node.full_name,
                node.manufacturer,
                node.supplier,
                node.manufacture_date,
                node.install_date,
                passport.commissioned_at if depth == 0 else None,
            )
            fill = (
                "DCE5F3"
                if depth == 0
                else (
                    "DDE5FA"
                    if depth == 1
                    else (
                        "EFF3F8"
                        if has_children
                        else "FAFBFD" if index % 2 == 0 else "FFFFFF"
                    )
                )
            )
            color = "172B4D" if depth == 0 else "293F78" if depth == 1 else "253247"
            lines = max(
                (
                    max(
                        (
                            math.ceil(len(part) / capacities[column])
                            for part in value.splitlines()
                        ),
                        default=1,
                    )
                    if isinstance(value, str)
                    else 1
                )
                for column, value in enumerate(values)
            )
            sheet.row_dimensions[row].height = max(
                30 if depth <= 1 else 26, lines * 14 + 10
            )
            sheet.row_dimensions[row].outlineLevel = min(depth, 7)
            for column, value in enumerate(values, 1):
                cell = sheet.cell(row, column)
                cls._text(cell, value)
                cell.font = Font(
                    name="Arial",
                    size=10,
                    color=color,
                    bold=depth <= 1 or (column == 1 and has_children),
                )
                cell.fill = PatternFill("solid", fgColor=fill)
                cell.alignment = Alignment(
                    horizontal="center" if column >= 8 else "left",
                    vertical="center",
                    wrap_text=True,
                )
                cell.number_format = "dd.mm.yyyy" if column >= 8 else "@"
                if depth == 1:
                    cell.border = Border(
                        top=Side(style="thin", color="AABBD8"),
                        bottom=Side(style="thin", color="AABBD8"),
                    )
        stream = io.BytesIO()
        workbook.save(stream)
        stream.seek(0)
        return stream

    @classmethod
    async def export(cls, session: AsyncSession, passport_id: int):
        passport = await ProductPassportService.get_by_id(session, passport_id)
        if passport is None:
            return None
        data = ProductPassportService.to_response(passport)
        stream = await asyncio.to_thread(cls.generate_excel, data)
        filename = f"Паспорт_{data.locomotive_model_name}_{data.product_number}.xlsx"
        filename = re.sub(r'[\\/:*?"<>|\r\n]', "_", filename)
        return stream, filename
